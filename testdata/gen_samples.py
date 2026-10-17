"""Generate sample decks used by tests and demos.

Requires: pip install python-pptx pillow
Usage: python3 testdata/gen_samples.py
"""
import io
import os
import re
import zipfile

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
NAVY = RGBColor(0x0F, 0x34, 0x60)
PINK = RGBColor(0xE9, 0x45, 0x60)
INK = RGBColor(0x1A, 0x1A, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x8A, 0x8F, 0xA3)


def png(w, h, color, text=None):
    img = Image.new("RGB", (w, h), color)
    d = ImageDraw.Draw(img)
    if text:
        d.text((10, h // 2 - 6), text, fill=(255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


LOGO = png(200, 80, (233, 69, 96), "ACME").getvalue()


def add_chrome(prs, slide, n):
    sw, sh = prs.slide_width, prs.slide_height
    slide.shapes.add_picture(io.BytesIO(LOGO), int(sw * 0.86), int(sh * 0.89), int(sw * 0.10), int(sh * 0.07))
    tb = slide.shapes.add_textbox(int(sw * 0.05), int(sh * 0.91), int(sw * 0.1), int(sh * 0.05))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = str(n)
    r.font.size = Pt(11)
    r.font.color.rgb = MUTED


def style_title(shape, color=INK, size=36):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = color


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def acme():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    sw, sh = prs.slide_width, prs.slide_height
    L_TITLE, L_CONTENT, L_TITLE_ONLY = prs.slide_layouts[0], prs.slide_layouts[1], prs.slide_layouts[5]
    n = 0

    def title_slide(title, sub):
        nonlocal n
        n += 1
        s = prs.slides.add_slide(L_TITLE)
        bg(s, NAVY)
        s.shapes.title.text = title
        s.shapes.title.left, s.shapes.title.top = int(sw * 0.06), int(sh * 0.36)
        s.shapes.title.width, s.shapes.title.height = int(sw * 0.7), int(sh * 0.16)
        style_title(s.shapes.title, WHITE, 44)
        s.shapes.title.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        st = s.placeholders[1]
        st.text = sub
        st.left, st.top, st.width, st.height = int(sw * 0.06), int(sh * 0.54), int(sw * 0.6), int(sh * 0.08)
        for r in st.text_frame.paragraphs[0].runs:
            r.font.size = Pt(20)
            r.font.color.rgb = PINK
        st.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        add_chrome(prs, s, n)

    def text_image(title, bullets, color):
        nonlocal n
        n += 1
        s = prs.slides.add_slide(L_CONTENT)
        s.shapes.title.text = title
        t = s.shapes.title
        t.left, t.top, t.width, t.height = int(sw * 0.06), int(sh * 0.06), int(sw * 0.88), int(sh * 0.12)
        style_title(t)
        t.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        body = s.placeholders[1]
        body.left, body.top, body.width, body.height = int(sw * 0.06), int(sh * 0.22), int(sw * 0.42), int(sh * 0.64)
        tf = body.text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(18)
                r.font.color.rgb = INK
        s.shapes.add_picture(png(640, 480, color), int(sw * 0.52), int(sh * 0.22), int(sw * 0.42), int(sh * 0.64))
        add_chrome(prs, s, n)

    def kpis(title, items):
        nonlocal n
        n += 1
        s = prs.slides.add_slide(L_TITLE_ONLY)
        s.shapes.title.text = title
        t = s.shapes.title
        t.left, t.top, t.width, t.height = int(sw * 0.06), int(sh * 0.06), int(sw * 0.88), int(sh * 0.12)
        style_title(t)
        t.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        k = len(items)
        gap = 0.02
        w = (0.88 - gap * (k - 1)) / k
        for i, (val, label) in enumerate(items):
            x = 0.06 + i * (w + gap)
            card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, int(sw * x), int(sh * 0.32), int(sw * w), int(sh * 0.3))
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(0xF4, 0xF5, 0xF9)
            card.line.fill.background()
            vb = s.shapes.add_textbox(int(sw * (x + 0.02)), int(sh * 0.36), int(sw * (w - 0.04)), int(sh * 0.12))
            r = vb.text_frame.paragraphs[0].add_run()
            r.text = val
            r.font.size, r.font.bold, r.font.color.rgb = Pt(40), True, PINK
            lb = s.shapes.add_textbox(int(sw * (x + 0.02)), int(sh * 0.50), int(sw * (w - 0.04)), int(sh * 0.08))
            r = lb.text_frame.paragraphs[0].add_run()
            r.text = label
            r.font.size, r.font.color.rgb = Pt(14), MUTED
        add_chrome(prs, s, n)

    def table(title, rows):
        nonlocal n
        n += 1
        s = prs.slides.add_slide(L_TITLE_ONLY)
        s.shapes.title.text = title
        t = s.shapes.title
        t.left, t.top, t.width, t.height = int(sw * 0.06), int(sh * 0.06), int(sw * 0.88), int(sh * 0.12)
        style_title(t)
        t.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        gf = s.shapes.add_table(len(rows), len(rows[0]), int(sw * 0.06), int(sh * 0.22), int(sw * 0.88), int(sh * 0.5))
        tbl = gf.table
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = tbl.cell(i, j)
                c.text = val
                c.fill.solid()
                c.fill.fore_color.rgb = NAVY if i == 0 else (RGBColor(0xEE, 0xF1, 0xF7) if i % 2 else WHITE)
                for r in c.text_frame.paragraphs[0].runs:
                    r.font.size = Pt(14)
                    r.font.bold = i == 0
                    r.font.color.rgb = WHITE if i == 0 else INK
        add_chrome(prs, s, n)

    title_slide("Acme Growth Strategy 2026", "Board review - Q3 update")
    text_image("Why now", ["Market grew 38% YoY", "Competitors consolidating", "Our CAC dropped by half"], (15, 52, 96))
    kpis("Q3 at a glance", [("42%", "Revenue growth"), ("1.2M", "Active users"), ("98.9%", "Uptime")])
    text_image("Product roadmap", ["Self-serve onboarding", "Enterprise SSO", "Usage-based pricing", "EU data residency"], (233, 69, 96))
    table("Regional performance", [["Region", "Revenue", "Growth", "NPS"], ["North America", "$12.4M", "+31%", "62"], ["Europe", "$8.1M", "+44%", "58"], ["APAC", "$4.9M", "+52%", "55"]])
    kpis("Unit economics", [("$38", "CAC"), ("$410", "LTV"), ("10.8x", "LTV / CAC"), ("7 mo", "Payback")])
    text_image("Go-to-market", ["Partner channel in EMEA", "Vertical playbooks", "Customer advisory board"], (90, 90, 120))
    table("Hiring plan", [["Team", "Today", "Q4", "Q1"], ["Engineering", "48", "56", "64"], ["Sales", "22", "28", "34"]])
    title_slide("Thank you", "questions@acme.example")

    buf = io.BytesIO()
    prs.save(buf)
    return retheme(buf.getvalue())


def retheme(data):
    """Swap the default Office theme for a branded color/font scheme."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    dst = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    colors = {"dk1": "1A1A2E", "lt1": "FFFFFF", "dk2": "0F3460", "lt2": "EEF1F7", "accent1": "E94560",
              "accent2": "0F3460", "accent3": "16C79A", "accent4": "F5A623", "accent5": "533483", "accent6": "8A8FA3"}
    for item in src.infolist():
        b = src.read(item.filename)
        if item.filename.startswith("ppt/theme/theme"):
            s = b.decode("utf-8")
            for name, hexv in colors.items():
                s = re.sub(r"(<a:%s>)(.*?)(</a:%s>)" % (name, name), r'\1<a:srgbClr val="%s"/>\3' % hexv, s, flags=re.S)
            s = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*', r"\1Montserrat", s)
            s = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*', r"\1Inter", s)
            b = s.encode("utf-8")
        dst.writestr(item, b)
    dst.close()
    return out.getvalue()


if __name__ == "__main__":
    with open(os.path.join(HERE, "acme.pptx"), "wb") as f:
        f.write(acme())
    print("wrote testdata/acme.pptx")
